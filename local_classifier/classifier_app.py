#!/usr/bin/env python3
"""Local document classifier + field extractor (offline-first).

Classifies a document (ktp/kk/slip/mutasi/sk) and extracts fields. OCR strategy:
  1. Digital PDF text via pypdfium2 (fast, exact) for slip/mutasi.
  2. For scans/photos: corp PaddleOCR if PADDLE_OCR_URL is set (best quality),
     else local Tesseract on a binarized image.

Run from the repo root with the shared venv:
  $env:Path = "C:\\Program Files\\Tesseract-OCR;" + $env:Path
  .venv\\Scripts\\python.exe -m uvicorn classifier_app:app \
      --app-dir local_classifier --host 127.0.0.1 --port 8020
"""
from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile
from collections import defaultdict

from fastapi import FastAPI, File, HTTPException, UploadFile

import pypdfium2 as pdfium
from PIL import Image, ImageOps

# PaddleOCR (local, optional) — disable oneDNN to dodge a paddlepaddle 3.x CPU
# bug, and skip the model-source connectivity check for faster init.
os.environ.setdefault("FLAGS_use_mkldnn", "0")
os.environ.setdefault("PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK", "True")

_PADDLE: dict = {}      # lazy PaddleOCR instances, cached by recognizer mode

app = FastAPI(title="Local Document Classifier + Extractor", version="2.0.0")

KEYWORDS: dict[str, list[tuple[str, int]]] = {
    "ktp": [
        ("kartu tanda penduduk", 5), ("nik", 3), ("kewarganegaraan", 2),
        ("golongan darah", 2), ("gol. darah", 2), ("tempat/tgl lahir", 2),
        ("status perkawinan", 2), ("provinsi", 1), ("kabupaten", 1), ("agama", 1),
    ],
    "kk": [
        ("kartu keluarga", 5), ("nomor kk", 3), ("no. kk", 3),
        ("kepala keluarga", 3), ("hubungan dalam keluarga", 3),
        ("hubungan keluarga", 2), ("nama lengkap", 1), ("tempat lahir", 1),
    ],
    "slip": [
        ("slip gaji", 5), ("gaji pokok", 4), ("take home pay", 3),
        ("total diterima", 3), ("penghasilan", 2), ("tunjangan", 2),
        ("potongan", 2), ("upah", 2), ("thp", 2), ("periode", 1),
    ],
    "mutasi": [
        ("mutasi rekening", 5), ("rekening koran", 5), ("account statement", 4),
        ("e-statement", 4), ("saldo akhir", 3), ("tanggal transaksi", 2),
        ("uraian transaksi", 3), ("mutasi", 1), ("saldo", 1), ("debet", 1),
        ("kredit", 1), ("no. rekening", 1), ("nomor rekening", 1),
    ],
    "sk": [
        ("surat keterangan kerja", 5), ("keterangan kerja", 4),
        ("surat keputusan", 4), ("dengan ini menerangkan", 3),
        ("karyawan tetap", 2), ("masa kerja", 2), ("diangkat", 2), ("jabatan", 1),
    ],
}


# ── OCR helpers ────────────────────────────────────────────────────────────

def _detect_tessdata() -> tuple[str | None, str]:
    """Default to system `eng` (proven on our KK/SK/slip/mutasi samples).
    Indonesian (ind) changed digit reading and regressed KK NIK extraction, so
    it is opt-in via TESS_USE_IND=1 (uses the bundled tessdata/ind.traineddata)."""
    if os.environ.get("TESS_USE_IND") == "1":
        local = os.path.join(os.path.dirname(__file__), "tessdata")
        if os.path.exists(os.path.join(local, "ind.traineddata")):
            return local, "ind+eng"
    return None, "eng"


_TESSDATA_DIR, _LANG = _detect_tessdata()


def _otsu(gray: "Image.Image") -> int:
    hist = gray.histogram()[:256]
    total = sum(hist)
    if total == 0:
        return 128
    sum_all = sum(i * hist[i] for i in range(256))
    sum_b = 0.0
    w_b = 0
    maximum = 0.0
    level = 128
    for i in range(256):
        w_b += hist[i]
        if w_b == 0:
            continue
        w_f = total - w_b
        if w_f == 0:
            break
        sum_b += i * hist[i]
        m_b = sum_b / w_b
        m_f = (sum_all - sum_b) / w_f
        between = w_b * w_f * (m_b - m_f) ** 2
        if between >= maximum:
            level = i
            maximum = between
    return level


def _preprocess(img: "Image.Image") -> "Image.Image":
    """Grayscale + autocontrast + upscale + Otsu binarize — helps Tesseract."""
    g = ImageOps.autocontrast(img.convert("L"))
    w, h = g.size
    m = max(w, h)
    if m < 2000:
        s = 2000 / m
        g = g.resize((int(w * s), int(h * s)))
    t = _otsu(g)
    return g.point(lambda x: 255 if x > t else 0)


def _run_tess(img: "Image.Image", psm: str | None = None) -> str:
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tf:
        img.save(tf, "PNG")
        tmp = tf.name
    try:
        cmd = ["tesseract", tmp, "stdout", "-l", _LANG]
        if _TESSDATA_DIR:
            cmd += ["--tessdata-dir", _TESSDATA_DIR]
        if psm:
            cmd += ["--psm", psm]
        return subprocess.run(cmd, capture_output=True, text=True).stdout or ""
    except Exception:
        return ""
    finally:
        os.unlink(tmp)


def _ocr_ktp_aggressive(data: bytes, filename: str) -> str:
    """Hard KTP scans (small ID on a mostly-blank page, low contrast): render
    hi-res, try full image + top-left crop, boost contrast, multiple PSMs; keep
    the densest text. Best-effort — NIK on poor scans may still be unreadable."""
    if not shutil.which("tesseract"):
        return ""
    from PIL import ImageEnhance

    is_pdf = data[:5] == b"%PDF-" or filename.lower().endswith(".pdf")
    try:
        if is_pdf:
            pdf = pdfium.PdfDocument(data)
            pil = pdf[0].render(scale=3.0).to_pil()
            pdf.close()
        else:
            pil = Image.open(io.BytesIO(data))
    except Exception:
        return ""
    g = pil.convert("L")
    W, H = g.size
    # Full page + the top-left corner where scanned IDs usually sit.
    crops = [g, g.crop((0, 0, int(0.65 * W), int(0.33 * H)))]
    best = ""

    def score(s: str) -> int:
        return len(re.findall(r"[A-Za-z0-9]", s))

    for c in crops:
        w, h = c.size
        s = 2800 / max(w, h)
        if s > 1:
            c = c.resize((int(w * s), int(h * s)))
        c = ImageEnhance.Contrast(ImageOps.autocontrast(c)).enhance(2.0)
        for psm in ("6", "4", "11"):
            out = _run_tess(c, psm=psm)
            if score(out) > score(best):
                best = out
    return best


# ── Local PaddleOCR (no corp host) — reads ID cards far better than Tesseract ─

def _get_paddle(mode: str = "server"):
    """Lazy PaddleOCR. mode='server' = accurate recognizer (KTP names/fields);
    mode='mobile' = fast recognizer (KK table — many boxes; NIK digits stay
    accurate). Both share the fast mobile detector."""
    if mode in _PADDLE:
        return _PADDLE[mode]
    rec = "PP-OCRv5_server_rec" if mode == "server" else "PP-OCRv5_mobile_rec"
    try:
        from paddleocr import PaddleOCR
        _PADDLE[mode] = PaddleOCR(
            lang="en",
            text_detection_model_name="PP-OCRv5_mobile_det",
            text_recognition_model_name=rec,
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=False,
            enable_mkldnn=False,
        )
    except Exception:
        _PADDLE[mode] = None
    return _PADDLE[mode]


def _paddle_fragments(pil: "Image.Image", mode: str = "server") -> list[str]:
    """Run local PaddleOCR on a PIL image → ordered list of text fragments."""
    ocr = _get_paddle(mode)
    if ocr is None:
        return []
    try:
        import numpy as np
        frags: list[str] = []
        for r in ocr.predict(np.array(pil.convert("RGB"))):
            frags.extend(r.get("rec_texts", []))
        return [f for f in frags if f and f.strip()]
    except Exception:
        return []


def paddle_available() -> bool:
    return _get_paddle("mobile") is not None


def _smart_crop(pil: "Image.Image") -> "Image.Image":
    """If the document is a small ID scanned on a mostly-blank page, crop to the
    dark-content bounding box and upscale so PaddleOCR sees larger text. A
    full-frame photo (content fills the page) is returned unchanged."""
    g = pil.convert("L")
    W, H = g.size
    mask = g.point(lambda x: 255 if x < 165 else 0)  # reasonably dark = real content
    bbox = mask.getbbox()
    if not bbox:
        return pil
    bw, bh = bbox[2] - bbox[0], bbox[3] - bbox[1]
    if bw >= 0.88 * W and bh >= 0.88 * H:
        return pil  # content fills the page already
    pad = 24
    box = (max(0, bbox[0] - pad), max(0, bbox[1] - pad),
           min(W, bbox[2] + pad), min(H, bbox[3] + pad))
    c = pil.crop(box)
    cw, ch = c.size
    s = 2800 / max(cw, ch)
    if s > 1:
        c = c.resize((int(cw * s), int(ch * s)))
    return c


def _render_for_paddle(data: bytes, filename: str, scale: float) -> "Image.Image | None":
    is_pdf = data[:5] == b"%PDF-" or filename.lower().endswith(".pdf")
    try:
        if is_pdf:
            pdf = pdfium.PdfDocument(data)
            pil = pdf[0].render(scale=scale).to_pil()
            pdf.close()
            return pil
        return Image.open(io.BytesIO(data))
    except Exception:
        return None


def _paddle_ktp_fragments(data: bytes, filename: str) -> list[str]:
    """PaddleOCR a KTP — auto-crop to content + upscale, then a single predict."""
    pil = _render_for_paddle(data, filename, 3.0)
    return _paddle_fragments(_smart_crop(pil)) if pil else []


def _paddle_items(pil: "Image.Image", mode: str = "mobile") -> list[dict]:
    """Like _paddle_fragments but keeps each box's position (x, vertical centre,
    height) so table rows can be reconstructed."""
    ocr = _get_paddle(mode)
    if ocr is None:
        return []
    try:
        import numpy as np
        items = []
        for r in ocr.predict(np.array(pil.convert("RGB"))):
            texts = r.get("rec_texts", []) or []
            boxes = r.get("rec_boxes", None)
            if boxes is None:
                boxes = []
                for p in (r.get("dt_polys", []) or []):
                    xs = [float(pt[0]) for pt in p]
                    ys = [float(pt[1]) for pt in p]
                    boxes.append([min(xs), min(ys), max(xs), max(ys)])
            for t, b in zip(texts, boxes):
                if t and str(t).strip():
                    y1, y2 = float(b[1]), float(b[3])
                    items.append({"text": str(t).strip(), "x": float(b[0]),
                                  "yc": (y1 + y2) / 2, "h": max(1.0, y2 - y1)})
        return items
    except Exception:
        return []


_KK_NON_NAME = re.compile(
    r"provinsi|kota|kabupaten|kecamatan|kelurahan|desa|alamat|kode|pos|nomor|"
    r"\bkk\b|tanggal|dikeluarkan|kepala|status|\bnama\b|\bnik\b|jenis|kelamin|"
    r"agama|pekerjaan|pendidikan|gol|darah|hubungan|kewarganegaraan|tempat|"
    r"lahir|kawin|tercatat|\brt\b|\brw\b|islam|laki|perempuan",
    re.IGNORECASE,
)


def extract_kk_from_items(items: list[dict]) -> dict:
    """Reconstruct KK fields from positioned PaddleOCR boxes: the top-most
    16-digit = Nomor KK; every other 16-digit NIK is paired with the name on
    the same table row (same vertical centre, to its left)."""
    fields = {"nomorKK": None, "kepalaKeluarga": None, "alamat": None, "members": []}
    niks = sorted([it for it in items if re.fullmatch(r"\d{16}", it["text"])], key=lambda it: it["yc"])

    txt = "\n".join(it["text"] for it in sorted(items, key=lambda it: (it["yc"], it["x"])))
    # Nomor KK: the labelled 16-digit at the top; else the top-most 16-digit.
    fields["nomorKK"] = _first(r"No\.?\s*[:.]?\s*(\d{16})", txt) or (niks[0]["text"] if niks else None)

    def is_name(t: str) -> bool:
        words = t.split()
        letters = re.findall(r"[A-Za-z]", t)
        uppers = sum(1 for c in t if c.isalpha() and c.isupper())
        return (2 <= len(words) <= 5 and len(t) <= 45 and len(letters) >= 4
                and uppers >= 0.7 * len(letters) and not _KK_NON_NAME.search(t))

    # Member NIKs = all NIKs except the one equal to Nomor KK.
    member_niks = [n for n in niks if n["text"] != fields["nomorKK"]]
    members, seen = [], set()
    for nik in member_niks:
        if nik["text"] in seen:
            continue
        seen.add(nik["text"])
        row = [it for it in items
               if abs(it["yc"] - nik["yc"]) < max(25.0, 0.8 * nik["h"])
               and it["x"] < nik["x"] and is_name(it["text"])]
        row.sort(key=lambda it: it["x"])
        if row:
            members.append({"nama": re.sub(r"\s+", " ", row[0]["text"]).strip(),
                            "nik": nik["text"], "hubungan": ""})
    fields["members"] = members[:12]
    if members:
        fields["kepalaKeluarga"] = members[0]["nama"]

    fields["alamat"] = _first(r"Alamat\s*[:.]?\s*([^\n]+)", txt)
    return fields


def extract_ktp_from_fragments(frags: list[str]) -> dict:
    """The server recognizer returns clean, labelled KTP text (one fragment per
    field/value), so the label-based extractor works once fragments are joined.
    The first card on the page (the applicant) wins via re.search."""
    return extract_ktp_fields("\n".join(frags))


def _ocr_pil(pil: "Image.Image") -> str:
    """OCR the rendered image as-is first (works well for KK). Only when that
    yields almost nothing (e.g. a hard KTP photo) fall back to grayscale +
    upscale + Otsu binarize."""
    if not shutil.which("tesseract"):
        return ""
    text = _run_tess(pil)
    if len(re.sub(r"\s", "", text)) >= 40:
        return text
    g = pil.convert("L")
    w, h = g.size
    m = max(w, h)
    if m < 2200:
        s = 2200 / m
        g = g.resize((int(w * s), int(h * s)))
    alt = _run_tess(g)
    if len(re.sub(r"\s", "", alt)) < 40:
        alt2 = _run_tess(g.point(lambda x: 255 if x > _otsu(g) else 0))
        if len(alt2.strip()) > len(alt.strip()):
            alt = alt2
    return alt if len(alt.strip()) > len(text.strip()) else text


def _ocr_pages(pdf: "pdfium.PdfDocument", max_pages: int = 4) -> str:
    parts = []
    for i in range(min(len(pdf), max_pages)):
        try:
            parts.append(_ocr_pil(pdf[i].render(scale=2.0).to_pil()))
        except Exception:
            continue
    return "\n".join(parts)


def _paddle_ocr(data: bytes, filename: str) -> str | None:
    """Corp PaddleOCR markdown service — used only when PADDLE_OCR_URL is set.

    Set env: PADDLE_OCR_URL (e.g. http://10.213.128.80:8090/predict/markdown),
    PADDLE_OCR_KEY (X-API-Key). Falls back to Tesseract when unset/unreachable.
    """
    url = os.environ.get("PADDLE_OCR_URL")
    if not url:
        return None
    try:
        import requests  # lazy

        headers = {}
        key = os.environ.get("PADDLE_OCR_KEY") or os.environ.get("OCR_API_KEY")
        if key:
            headers["X-API-Key"] = key
        params = {}
        skip = os.environ.get("PADDLE_SKIP_ORIENTATION")
        if skip:
            params["skip_orientation"] = skip
        resp = requests.post(
            url, params=params, files={"file": (filename or "doc", data)},
            headers=headers, timeout=float(os.environ.get("PADDLE_TIMEOUT", "60")),
        )
        if resp.status_code != 200:
            return None
        j = resp.json()
        parts = []
        for res in ((j.get("data") or {}).get("json_result") or []):
            for blk in (res.get("parsing_res_list") or []):
                c = blk.get("block_content")
                if c:
                    parts.append(str(c))
        text = "\n".join(parts)
        return text if len(re.sub(r"\s", "", text)) >= 20 else None
    except Exception:
        return None


def extract_text(data: bytes, filename: str = "") -> str:
    """Digital text first; PaddleOCR (if configured) then Tesseract for scans."""
    is_pdf = data[:5] == b"%PDF-" or filename.lower().endswith(".pdf")
    if is_pdf:
        try:
            pdf = pdfium.PdfDocument(data)
        except Exception:
            return ""
        try:
            digital = []
            for i in range(len(pdf)):
                digital.append(pdf[i].get_textpage().get_text_range() or "")
            text = "\n".join(digital)
            if len(re.sub(r"\s", "", text)) >= 40:
                return text  # digital PDF (slip / mutasi)
            paddle = _paddle_ocr(data, filename)
            return text + "\n" + (paddle if paddle else _ocr_pages(pdf))
        finally:
            pdf.close()
    paddle = _paddle_ocr(data, filename)
    if paddle:
        return paddle
    try:
        return _ocr_pil(Image.open(io.BytesIO(data)).convert("RGB"))
    except Exception:
        return ""


# ── Classification ─────────────────────────────────────────────────────────

def classify_text(text: str) -> tuple[str, str, dict[str, int]]:
    low = text.casefold()
    scores = {label: sum(w for kw, w in kws if kw in low) for label, kws in KEYWORDS.items()}
    best = max(scores, key=lambda k: scores[k])
    best_score = scores[best]
    if best_score < 3:
        return "unknown", "low", scores
    conf = "high" if best_score >= 7 else "medium" if best_score >= 5 else "low"
    return best, conf, scores


def _classify_bytes(filename: str, data: bytes) -> dict:
    if not data:
        return {"filename": filename, "document_type": "unknown", "confidence": "low", "scores": {}, "error": "empty"}
    try:
        text = extract_text(data, filename)
        frags = None
        # A KTP photo is unreadable by Tesseract (≈0 chars). Run ONE local
        # PaddleOCR pass and reuse it for BOTH classification and (if it is a
        # KTP) field extraction — so the KTP is never OCR'd twice.
        if len(re.sub(r"\s", "", text)) < 30:
            frags = _paddle_ktp_fragments(data, filename)
            if frags:
                text += "\n" + "\n".join(frags)
        doc_type, conf, scores = classify_text(text)
        result = {"filename": filename, "document_type": doc_type, "confidence": conf, "scores": scores}
        if doc_type == "ktp" and frags:
            result["fields"] = extract_ktp_from_fragments(frags)
        return result
    except Exception as exc:  # noqa: BLE001
        return {"filename": filename, "document_type": "unknown", "confidence": "low", "scores": {}, "error": str(exc)}


# ── Field extraction ───────────────────────────────────────────────────────

def _first(pattern: str, text: str):
    m = re.search(pattern, text, re.IGNORECASE)
    return m.group(1).strip() if m else None


def _money(s: str) -> float:
    """Parse Rupiah in either US (71,546.00) or Indonesian (10.000.000,00) style."""
    s = (s or "").replace("Rp", "").replace(" ", "").strip()
    if not s:
        return 0.0
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):       # comma is decimal (ID)
            s = s.replace(".", "").replace(",", ".")
        else:                                  # dot is decimal (US)
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", "") if re.search(r",\d{3}(\D|$)", s) else s.replace(",", ".")
    elif s.count(".") > 1 or re.search(r"\.\d{3}(\D|$)", s):
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _money_of(text: str, *patterns: str):
    for p in patterns:
        m = re.search(p, text, re.IGNORECASE)
        if m:
            return _money(m.group(1))
    return None


def extract_ktp_fields(text: str) -> dict:
    nik = _first(r"NIK\s*[:.]?\s*([0-9]{16})", text)
    if not nik:
        m = re.search(r"\b(\d{16})\b", text)
        nik = m.group(1) if m else None
    nama = _first(r"Nama\s*[:.]?\s*([A-Z][A-Za-z .'`-]{2,40})", text)
    gender = (
        "Laki-laki" if re.search(r"LAKI[- ]?LAKI", text, re.IGNORECASE)
        else "Perempuan" if re.search(r"PEREMPUAN", text, re.IGNORECASE) else None
    )
    ttl = _first(r"Tempat\s*/?\s*Tgl[. ]*Lahir\s*[:.]?\s*([^\n]+)", text)
    tempat, tgl = None, None
    if ttl:
        ttl = ttl.lstrip(":. ").strip()
        # Allow "SURABAYA, 11-04-1999" OR "SURABAYA11-04-1999" (no separator).
        m = re.search(r"([A-Za-z .'-]+?)[,\s]*(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})", ttl)
        if m:
            tempat, tgl = m.group(1).strip(" ,."), m.group(2).strip()
        else:
            tempat = ttl.strip()
    if not tgl:
        m = re.search(r"\b(\d{1,2}[-/]\d{1,2}[-/]\d{4})\b", text)
        tgl = m.group(1) if m else None
    alamat = _first(r"Alamat\s*[:.]?\s*([^\n]+)", text)
    status_kawin = (
        "Belum Kawin" if re.search(r"BELUM\s*KAWIN", text, re.IGNORECASE)
        else "Cerai" if re.search(r"CERAI", text, re.IGNORECASE)
        else "Kawin" if re.search(r"\bKAWIN\b", text, re.IGNORECASE) else None
    )
    return {"nik": nik, "nama": nama, "gender": gender,
            "tempatLahir": tempat, "tanggalLahir": tgl, "alamat": alamat,
            "statusPerkawinan": status_kawin}


def extract_kk_fields(text: str) -> dict:
    # Nomor KK is printed at the very top, just under "KARTU KELUARGA".
    head = text[:300]
    nomor = (_first(r"No\.?\s*[:.]?\s*([0-9OoBbEeIlSsZz]{14,18})", head)
             or _first(r"(?:Nomor|No)\.?\s*(?:KK)?\s*[:.]?\s*([0-9]{16})", text))
    if nomor:
        nomor = nomor.strip()
    # Case-SENSITIVE so we capture the all-uppercase name, not mixed-case noise.
    mk = (re.search(r"Nama Kepala Keluarga\s*[:.;]?\s*([A-Z][A-Z .'`-]{2,40})", text)
          or re.search(r"Kepala Keluarga\s*[:.;]?\s*([A-Z][A-Z .'`-]{2,40})", text))
    kepala = mk.group(1).strip() if mk else None
    if kepala:
        kepala = re.sub(r"\s+[A-Z]$", "", kepala).strip()
        if re.search(r"DINAS|KEPENDUDUKAN|PENCATATAN|CAMAT|LURAH", kepala, re.IGNORECASE):
            kepala = None
    alamat = _first(r"Alamat\s*[:.]?\s*([^\n]+)", text)
    # Members: UPPERCASE name followed by a 16-digit NIK (KK table rows).
    members = []
    seen = set()
    for m in re.finditer(r"([A-Z][A-Z .'`-]{3,40}?)[\s\[|:;]+(\d{16})\b", text):
        nama = re.sub(r"\s+", " ", m.group(1)).strip()
        nik = m.group(2)
        if nama.lower() in ("nama", "no", "no kk", "nomor kk", "kepala keluarga",
                            "nama kepala keluarga", "nama lengkap") or nik in seen:
            continue
        seen.add(nik)
        hub = "Kepala Keluarga" if kepala and nama.upper() == kepala.upper() else ""
        members.append({"nama": nama, "nik": nik, "hubungan": hub})
    return {"nomorKK": nomor, "kepalaKeluarga": kepala, "alamat": alamat, "members": members[:12]}


def extract_slip_fields(text: str) -> dict:
    # Tanggal pembayaran: BRI's "SLIP UPAH" prints "Periode Pembayaran" then the
    # date as DD.MM.YYYY on the next line, so grab the first date after that
    # header; else a labelled date; else a "Month YYYY".
    tgl = None
    head = re.search(r"Periode\s*Pembayaran", text, re.IGNORECASE)
    if head:
        md = re.search(r"(\d{1,2}[./-]\d{1,2}[./-]\d{2,4})", text[head.end(): head.end() + 160])
        if md:
            tgl = md.group(1)
    if not tgl:
        tgl = _first(
            r"(?:Tanggal\s*Pembayaran|Tgl\.?\s*Pembayaran|Tanggal\s*Bayar|Periode\s*Gaji|Period)\s*"
            r"[:.]?\s*([0-9]{1,2}[./ -][A-Za-z0-9]+[./ -][0-9]{2,4}|[A-Za-z]+\s+\d{4}|\d{1,2}\s+[A-Za-z]+\s+\d{4})",
            text,
        )
    upah = _money_of(
        text,
        r"Total\s*[-–]?\s*(?:Income|Upah|Penghasilan|Pendapatan|Bruto)\s*(?:/\s*\w+)?\s*[:.]?\s*(?:Rp)?\s*([\d.,]+)",
        r"(?:Penghasilan|Gaji)\s*Bruto\s*[:.]?\s*(?:Rp)?\s*([\d.,]+)",
        r"(?:Penghasilan)\s*[^\dA-Za-z]{0,6}([\d.,]{5,})",
    )
    potongan = _money_of(
        text,
        r"Total\s*[-–]?\s*(?:Deduction|Potongan|Pengurang)\s*(?:/\s*\w+)?\s*[:.]?\s*(?:Rp)?\s*([\d.,]+)",
        r"(?:Potongan|Pengurang|Deduction)\s*[^\dA-Za-z]{0,6}([\d.,]{3,})",
    )
    thr = _money_of(text, r"THR\s*[:.]?\s*(?:Rp)?\s*([\d.,]+)",
                    r"Tunjangan\s*Hari\s*Raya\s*[:.]?\s*(?:Rp)?\s*([\d.,]+)")
    bonus = _money_of(text, r"Bonus\s*[:.]?\s*(?:Rp)?\s*([\d.,]+)")
    thp = round(upah - potongan, 2) if (upah is not None and potongan is not None) else None
    if thp is None:
        thp = _money_of(
            text,
            r"(?:THP|Take\s*Home\s*Pay|Total\s*Diterima|Gaji\s*Bersih|Penghasilan\s*Bersih|Neto|Netto)\s*"
            r"[:.]?\s*(?:Rp)?\s*([\d.,]+)",
            r"(?:Amount\s*transfer\w*(?:\s*into)?|Jumlah\s*di\s*transfer)[^\d]{0,50}([\d.,]{5,})",
        )
    return {"tanggalPembayaran": tgl, "totalUpah": upah, "totalPotongan": potongan,
            "thp": thp, "thr": thr, "bonus": bonus}


def extract_sk_fields(text: str) -> dict:
    # Nomor surat / NOKEP — prefer NOKEP/Nomor (avoid the street "No. 44-46").
    nomor = (_first(r"NOKEP\s*[:.]?\s*([0-9.][\w./\-]{4,40})", text[:600])
             or _first(r"(?:No\.?\s*KEP|Nomor(?:\s*Surat)?|No\.?\s*Surat)\s*[:.]?\s*([0-9.][\w./\-]{4,40})", text[:600])
             or _first(r"\bNo\.?\s*[:.]?\s*([0-9][\w./\-]{6,40})", text[:400]))
    perusahaan = (_first(r"((?:PT|CV|UD|PD|PERUM)\.?\s+[A-Z][A-Za-z .'&-]{2,50}?)\s+menyatakan", text)
                  or _first(r"\b((?:PT|CV|UD|PD|PERUM)\.?\s+[A-Z][A-Za-z .'&-]{2,50})", text))
    # Nama — case-SENSITIVE so the lowercase "namanya" in BRI's boilerplate isn't
    # captured. "Sdr./Sdri. X" (BRI kutipan SK) or "Nama : X" / "Nama X" (SKK).
    mn = (re.search(r"Sdr\.?\s*/?\s*Sdri\.?\s+([A-Z][A-Za-z. '`-]{3,40})", text)
          or re.search(r"\bNama\b\s*[:.]?\s+([A-Z][A-Za-z. '`-]{3,40})", text))
    nama = re.sub(r"\s+", " ", mn.group(1)).strip() if mn else None
    # Jabatan — skip "Golongan Jabatan / Job Grade" style false hits.
    jabatan = None
    mj = re.search(r"Jabatan\s*(?:terkini\s*)?(?:sebagai\s*)?[:.]?\s+([A-Z][A-Za-z /&.,-]{3,50})", text)
    if mj and not re.search(r"job\s*grade|golongan|person\s*grade|unit\s*kerja", mj.group(1), re.IGNORECASE):
        jabatan = mj.group(1).strip()
    nik = _first(r"(?:No\.?\s*KTP|NIK)\s*[:.]?\s*([0-9]{16})", text)
    tgl_masuk = _first(
        r"(?:Terhitung\s*mulai\s*tanggal|Tanggal\s*Masuk|Mulai\s*(?:Bekerja|Kerja)|sejak\s*(?:tanggal)?)\s*[:.]?\s*"
        r"(\d{1,2}\s+[A-Za-z]+\s+\d{4}|\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4})",
        text,
    )
    status = (
        "Karyawan Tetap" if re.search(r"karyawan\s*tetap|pekerja\s*tetap", text, re.IGNORECASE)
        else "Kontrak (PKWT)" if re.search(r"\bkontrak\b|pkwt", text, re.IGNORECASE) else None
    )
    return {"perusahaan": perusahaan, "namaPekerja": nama, "jabatan": jabatan, "nik": nik,
            "tanggalMulai": tgl_masuk, "tanggalBerakhir": None,
            "statusKepegawaian": status, "masaKerja": None, "nomorSurat": nomor}


# ── Mutasi rekening (bank statement) ───────────────────────────────────────

def classify_txn(remark: str, kredit: float) -> str:
    r = remark.lower()
    # SAP-DD is the customer's payroll/salary transfer.
    if "sap-dd" in r or "sap dd" in r or "sapdd" in r or "sap_dd" in r:
        return "Gaji"
    if "thr" in r or "hari raya" in r:
        return "THR"
    if "bonus" in r:
        return "Bonus"
    if "cuti" in r:
        return "Tunjangan Cuti"
    if "tunjangan" in r:
        return "Tunjangan"
    if "gaji" in r or "payroll" in r or "salary" in r:
        return "Gaji"
    if "qris" in r:
        return "QRIS"
    if "top up" in r or "topup" in r:
        return "Top Up"
    if "penarikan" in r or "tarik tunai" in r or "atm" in r:
        return "Tarik Tunai"
    if "kartu kredit" in r:
        return "Bayar Kartu Kredit"
    if "briva" in r or "pembayaran" in r or "tagihan" in r:
        return "Pembayaran"
    if "transfer" in r or "from:" in r or "dari" in r or "ke " in r:
        return "Transfer Masuk" if kredit > 0 else "Transfer Keluar"
    return "Dana Masuk" if kredit > 0 else "Lainnya"


_TXN_RE = re.compile(
    r"(\d{2}/\d{2}/\d{2})\s+\d{2}:\d{2}:\d{2}\s+(.*?)\s+"
    r"([\d.,]+\.\d{2})\s+([\d.,]+\.\d{2})\s+([\d.,]+\.\d{2})"
    r"(?=\s+\d{2}/\d{2}/\d{2}\s+\d{2}:\d{2}:\d{2}|\s*$)",
    re.DOTALL,
)


def _parse_txns(text: str) -> list[dict]:
    # Strip BRImo's per-transaction footer ("Created By BRIMO\n<timestamp>\n
    # MSXC..._eStatement...") — its 4-digit-year timestamp doesn't match a
    # transaction date, so otherwise a description swallows the next rows.
    text = re.sub(
        r"Created By BRIMO.*?(?=\d{2}/\d{2}/\d{2}\s+\d{2}:\d{2}:\d{2}|\Z)",
        " ", text, flags=re.DOTALL,
    )
    # Strip the repeated per-page column header ("Tanggal Transaksi … Saldo
    # Balance") that otherwise splices across a page break.
    text = re.sub(r"Tanggal\s*Transaksi.*?Saldo\s*Balance", " ", text, flags=re.DOTALL)
    # Split at each transaction's "DD/MM/YY HH:MM:SS"; within the slice the first
    # three "….\d\d" amounts are Debet / Kredit / Saldo (account-header junk
    # after Saldo carries no decimal amount, so it is ignored). Robust to page
    # breaks that previously let one row swallow the next.
    out = []
    marks = list(re.finditer(r"(\d{2}/\d{2}/\d{2})\s+\d{2}:\d{2}:\d{2}", text))
    for i, mt in enumerate(marks):
        body = text[mt.end(): marks[i + 1].start() if i + 1 < len(marks) else len(text)]
        amts = re.findall(r"[\d.,]+\.\d{2}", body)
        if len(amts) < 3:
            continue
        debet = _money(amts[0])
        kredit = _money(amts[1])
        desc = re.sub(r"\s+", " ", body[: body.index(amts[0])]).strip()
        desc = re.sub(r"\s+\d{4,}$", "", desc).strip()  # drop trailing teller id
        out.append({
            "tanggal": mt.group(1), "remark": desc[:120],
            "nominal": kredit if kredit > 0 else debet,
            "dk": "Kredit" if kredit > 0 else "Debit", "_k": kredit,
        })
    return out


def _month(tgl: str) -> str:
    p = tgl.split("/")
    return f"{p[1]}/{p[2]}" if len(p) == 3 else tgl


def classify_mutasi(txns: list[dict]) -> tuple[list[dict], float | None]:
    """Gaji = the credit nominal recurring across the most distinct months."""
    months = defaultdict(set)
    for t in txns:
        if t["dk"] == "Kredit" and t["nominal"] > 0:
            months[t["nominal"]].add(_month(t["tanggal"]))
    gaji_nominal, best = None, 1
    for nom, mset in months.items():
        if len(mset) > best:
            best, gaji_nominal = len(mset), nom
    for t in txns:
        klas = classify_txn(t["remark"], t["_k"])
        if klas == "Dana Masuk" and gaji_nominal and t["nominal"] == gaji_nominal:
            klas = "Gaji"
        t["klasifikasi"] = klas
        t.pop("_k", None)
    return txns, gaji_nominal


def _txn_key(t: dict):
    p = t["tanggal"].split("/")
    return (p[2], p[1], p[0]) if len(p) == 3 else (t["tanggal"],)


# ── API ────────────────────────────────────────────────────────────────────

# ── SLIK OJK (read the parsed Excel: one row per credit facility) ────────────

_SLIK_XLSX = os.environ.get("SLIK_XLSX", r"C:\Users\BRI\Downloads\sample_slik_arie_parsing.xlsx")
_SLIK_ROWS: list[dict] = []


def _slik_num(s) -> float:
    try:
        return float(str(s).replace(",", "").strip() or 0)
    except (ValueError, TypeError):
        return 0.0


def _slik_months(mulai, tempo) -> int:
    try:
        ms = int(str(int(_slik_num(mulai)))[:6]); te = int(str(int(_slik_num(tempo)))[:6])
        return max(1, (te // 100 - ms // 100) * 12 + (te % 100 - ms % 100))
    except Exception:
        return 0


def _anuitas(principal: float, annual_rate: float, months: int) -> int:
    if months <= 0 or principal <= 0:
        return 0
    im = annual_rate / 12
    if im == 0:
        return round(principal / months)
    return round(principal * im / (1 - (1 + im) ** (-months)))


def _digits(v) -> str:
    return re.sub(r"\D", "", str(v)) if v is not None else ""


def _load_slik_rows() -> list[dict]:
    if _SLIK_ROWS:
        return _SLIK_ROWS
    try:
        import openpyxl
        rows = list(openpyxl.load_workbook(_SLIK_XLSX, read_only=True, data_only=True).active.iter_rows(values_only=True))
        hdr = [str(h).strip().lower() for h in rows[0]]
        for r in rows[1:]:
            if any(c is not None for c in r):
                _SLIK_ROWS.append(dict(zip(hdr, r)))
    except Exception:
        pass
    return _SLIK_ROWS


@app.get("/slik")
def slik(nik: str) -> dict:
    q = _digits(nik)
    # Excel stores the 16-digit NIK as a float and loses the last digits → match
    # on the first 13 digits.
    facs = [d for d in _load_slik_rows() if _digits(d.get("nik"))[:13] == q[:13]]
    if not facs:
        return {"ok": False, "error": f"NIK {nik} tidak ada di data SLIK"}
    loans = []
    total_aktif = 0
    worst = 1
    for d in facs:
        plafon = _slik_num(d.get("plafonawal") or d.get("plafon"))
        bunga = _slik_num(d.get("sukubunga"))  # already in percent
        months = _slik_months(d.get("tanggalmulai"), d.get("tanggaljatuhtempo"))
        angsuran = _anuitas(plafon, bunga / 100, months)
        aktif = int(_slik_num(d.get("kondisi"))) == 0  # 0 = aktif, 2 = lunas/non-aktif
        if aktif:
            total_aktif += angsuran
        kual = int(_slik_num(d.get("kolektibilitas")) or 1)
        worst = max(worst, kual)
        loans.append({
            "jenis": d.get("jeniskreditket"),
            "lembaga": d.get("ljkket"),
            "plafon": plafon,
            "baki": _slik_num(d.get("bakidebet")),
            "angsuran": angsuran,
            "sukuBunga": bunga,
            "tanggalMulai": str(d.get("tanggalmulai") or ""),
            "tanggalJatuhTempo": str(d.get("tanggaljatuhtempo") or ""),
            "aktif": aktif,
            "status": "Aktif" if aktif else (d.get("kondisiket") or "Non-Aktif"),
            "kualitas": kual,
        })
    return {
        "ok": True, "nik": nik, "namaDebitur": None,
        "loans": loans, "totalAngsuran": round(total_aktif),
        "kolekTerburuk": worst, "totalFasilitas": len(loans),
    }


@app.get("/health")
def health() -> dict:
    return {"status": "ok", "tesseract": bool(shutil.which("tesseract")),
            "lang": _LANG, "paddle": bool(os.environ.get("PADDLE_OCR_URL")),
            "slik": len(_load_slik_rows())}


@app.post("/classify")
async def classify(file: UploadFile = File(...)) -> dict:
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Empty upload")
    return _classify_bytes(file.filename or "document.pdf", data)


@app.post("/classify-batch")
async def classify_batch(files: list[UploadFile] = File(...)) -> dict:
    results = []
    for f in files:
        results.append(_classify_bytes(f.filename or "document.pdf", await f.read()))
    return {"count": len(results), "results": results}


@app.post("/extract-ktp")
async def extract_ktp(file: UploadFile = File(...)) -> dict:
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Empty upload")
    # Prefer local PaddleOCR — reads ID cards / NIK far better than Tesseract.
    frags = _paddle_ktp_fragments(data, file.filename or "ktp")
    if frags:
        pf = extract_ktp_from_fragments(frags)
        if pf.get("nik") or pf.get("nama"):
            return {"ok": True, "filename": file.filename, "fields": pf, "engine": "paddle"}
    # Fallback: Tesseract (digital text / aggressive OCR).
    text = extract_text(data, file.filename or "ktp")
    fields = extract_ktp_fields(text)
    if not fields.get("nik") or not fields.get("nama"):
        extra = _ocr_ktp_aggressive(data, file.filename or "ktp")
        if extra:
            f2 = extract_ktp_fields(text + "\n" + extra)
            fields = {k: (fields.get(k) or f2.get(k)) for k in f2}
    return {"ok": True, "filename": file.filename, "fields": fields, "engine": "tesseract"}


@app.post("/extract-kk")
async def extract_kk(file: UploadFile = File(...)) -> dict:
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Empty upload")
    # KK uses Tesseract: it reads this table correctly (kepala + members with the
    # right NIKs) and is fast. PaddleOCR's scattered table boxes scramble the
    # name↔NIK pairing, so it is intentionally NOT used here — it IS used for the
    # KTP photo, where Tesseract fails entirely.
    text = extract_text(data, file.filename or "kk")
    return {"ok": True, "filename": file.filename, "fields": extract_kk_fields(text), "engine": "tesseract"}


@app.post("/extract-slip")
async def extract_slip(files: list[UploadFile] = File(...)) -> dict:
    """One slip OR many (per payment date). Returns one record per file."""
    records = []
    for f in files:
        data = await f.read()
        if not data:
            continue
        rec = extract_slip_fields(extract_text(data, f.filename or "slip"))
        rec["fileName"] = f.filename
        records.append(rec)
    return {"ok": True, "records": records}


@app.post("/extract-sk")
async def extract_sk(file: UploadFile = File(...)) -> dict:
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Empty upload")
    return {"ok": True, "filename": file.filename, "fields": extract_sk_fields(extract_text(data, file.filename or "sk"))}


@app.post("/extract-mutasi")
async def extract_mutasi(files: list[UploadFile] = File(...)) -> dict:
    """Accepts one OR many statement files (e.g. 6 months) and merges them."""
    all_txns: list[dict] = []
    no_rek = None
    names = []
    for f in files:
        data = await f.read()
        if not data:
            continue
        names.append(f.filename)
        text = extract_text(data, f.filename or "mutasi")
        all_txns.extend(_parse_txns(text))
        if not no_rek:
            no_rek = _first(r"(?:No\.?\s*Rekening|Account No)\s*[:.]?\s*(\d{6,})", text)
    all_txns.sort(key=_txn_key)
    txns, gaji_nominal = classify_mutasi(all_txns)

    def total(klas: str) -> float:
        return round(sum(t["nominal"] for t in txns if t["klasifikasi"] == klas and t["dk"] == "Kredit"), 2)

    return {
        "ok": True, "filename": names,
        "fields": {
            "transactions": txns[:600], "noRekening": no_rek, "count": len(txns),
            "totalKredit": round(sum(t["nominal"] for t in txns if t["dk"] == "Kredit"), 2),
            "totalDebet": round(sum(t["nominal"] for t in txns if t["dk"] == "Debit"), 2),
            "gajiNominal": gaji_nominal,
            "ringkasan": {
                "Gaji": total("Gaji"), "THR": total("THR"), "Bonus": total("Bonus"),
                "Tunjangan": total("Tunjangan"), "Tunjangan Cuti": total("Tunjangan Cuti"),
            },
        },
    }
